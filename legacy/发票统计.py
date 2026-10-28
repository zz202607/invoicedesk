#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
发票统计工具 - 接收命令行参数
"""

import os
import sys
import io
import json
import re

# 设置标准输出编码
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
sys.stderr = io.TextIOWrapper(sys.stderr.buffer, encoding='utf-8', errors='replace')

try:
    import fitz  # PyMuPDF
except ImportError:
    print("错误: 未安装 PyMuPDF")
    sys.exit(1)

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
except ImportError:
    print("错误: 未安装 openpyxl")
    sys.exit(1)


def extract_text_from_pdf(pdf_path):
    """使用 PyMuPDF 提取 PDF 文本"""
    try:
        doc = fitz.open(pdf_path)
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        return text
    except Exception as e:
        return f"[解析错误: {str(e)}]"


def extract_amount(text):
    """从发票文本中提取价税合计金额（小写金额）"""
    if not text or len(text) < 10:
        return 0

    # 最优先：匹配"（小写）"后面的金额 —— 这就是价税合计
    patterns = [
        r'[（(]小写[）)]\s*[¥￥]\s*([\d,]+\.\d{2})',
        r'小写[）)]\s*[¥￥]\s*([\d,]+\.\d{2})',
        r'[（(]小写[）)]\s*([\d,]+\.\d{2})',
    ]
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            amount_str = match.group(1).replace(',', '')
            try:
                amount = float(amount_str)
                if 0 < amount < 100000000:
                    return amount
            except ValueError:
                continue

    # 其次：找"价税合计"后面最近的金额
    price_total_idx = text.find('价税合计')
    if price_total_idx >= 0:
        after_text = text[price_total_idx:]
        # 找"（小写）"或"小写"后的¥金额
        m = re.search(r'小写[）)]?\s*[¥￥]?\s*([\d,]+\.\d{2})', after_text)
        if m:
            try:
                amount = float(m.group(1).replace(',', ''))
                if 0 < amount < 100000000:
                    return amount
            except ValueError:
                pass
        # 没有小写标记，取价税合计后最大的金额（通常大写后面就是小写）
        amounts_after = re.findall(r'[¥￥]\s*([\d,]+\.\d{2})', after_text)
        if amounts_after:
            try:
                vals = [float(a.replace(',', '')) for a in amounts_after]
                vals = [v for v in vals if 0 < v < 100000000]
                if vals:
                    return max(vals)
            except ValueError:
                pass

    # 备用：文本中所有 ¥ 金额，取最大的（价税合计 > 不含税金额）
    all_amounts = re.findall(r'[¥￥]\s*([\d,]+\.\d{2})', text)
    if all_amounts:
        valid = []
        for a in all_amounts:
            try:
                val = float(a.replace(',', ''))
                if 1 < val < 100000000:
                    valid.append(val)
            except ValueError:
                pass
        if valid:
            return max(valid)

    return 0


def extract_item_name(text):
    """提取货物/服务名称"""
    if not text:
        return ''
    
    patterns = [
        r'\*([^*]+)\*\s*([^\s\n\r]{2,50})',
        r'货物或应税劳务[\s\S]*?名称[\s:：]*([^\s\n\r]{2,50})',
        r'品名[\s:：]*([^\s\n\r]{2,50})',
    ]
    
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            name = (match.group(2) if match.lastindex and match.lastindex >= 2 else match.group(1)).strip()
            if 1 < len(name) < 100:
                return name
    
    return ''


def extract_invoice_number(text):
    """提取发票号码"""
    if not text:
        return ''
    
    patterns = [
        r'发票号码[\s:：]*(\d{8,20})',
        r'No[\s.:]*(\d{8,20})',
    ]
    
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    
    return ''


def extract_date(text):
    """提取开票日期"""
    if not text:
        return ''
    
    patterns = [
        r'(\d{4})年(\d{1,2})月(\d{1,2})日',
        r'(\d{4})-(\d{2})-(\d{2})',
    ]
    
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            if '年' in match.group(0):
                try:
                    return f"{match.group(1)}-{int(match.group(2)):02d}-{int(match.group(3)):02d}"
                except (ValueError, IndexError):
                    pass
            return match.group(1)
    
    return ''


def extract_seller(text):
    """提取销售方名称"""
    if not text:
        return ''
    
    patterns = [
        r'销售方[\s\S]*?名称[\s:：]*([^\s税号地址]{2,50}?)(?:\s{2,}|纳税人识别号|地址|开户行|$)',
    ]
    
    for pattern in patterns:
        match = re.search(pattern, text)
        if match:
            return match.group(1).strip()
    
    return ''


CATEGORIES = {
    "差旅费": ["住宿", "酒店", "宾馆", "差旅", "机票", "火车票"],
    "餐饮费": ["餐饮", "餐费"],
    "汽车费用": ["加油", "油费", "停车费", "过路费", "过桥费", "高速公路", "通行费", "ETC"],
    "福利费": ["化妆品", "粮油", "服装", "鞋帽", "药品", "福利"],
    "业务招待费": ["礼品", "烟酒", "茶叶", "奢侈", "娱乐", "旅游"],
    "办公费用": ["办公", "文具", "打印", "纸张", "电脑", "手机", "电信", "网络", "快递", "桶装水", "顺丰"],
    "广告费用": ["广告", "宣传", "会务", "标签", "瓶身"],
    "培训费": ["培训", "学费", "教育"],
    "材料设备": ["材料", "原料", "设备", "泵", "膜", "五金", "工具", "配件", "机械", "仪器", "超滤", "热缩", "风机", "吨桶", "硅藻土", "制冷机", "硅"],
}


def categorize(item_name, file_name, full_text):
    """分类"""
    text = f"{item_name} {file_name} {full_text}".lower()
    
    for category, keywords in CATEGORIES.items():
        for kw in keywords:
            if kw.lower() in text:
                return category
    
    return "其他费用"


def parse_invoice(pdf_path):
    """解析单个PDF发票"""
    file_name = os.path.basename(pdf_path)
    text = extract_text_from_pdf(pdf_path)
    
    result = {
        "filePath": pdf_path,
        "fileName": file_name,
        "textLength": len(text),
        "amount": 0,
        "itemName": "",
        "invoiceNumber": "",
        "date": "",
        "seller": "",
        "category": "",
        "parseStatus": "未提取"
    }
    
    if text and len(text) > 10 and "[解析错误" not in text:
        result["amount"] = extract_amount(text)
        result["itemName"] = extract_item_name(text)
        result["invoiceNumber"] = extract_invoice_number(text)
        result["date"] = extract_date(text)
        result["seller"] = extract_seller(text)
        result["parseStatus"] = "成功" if result["amount"] > 0 else "未提取到金额"
    
    # 备用：从文件名提取
    if result["amount"] == 0:
        name_match = re.search(r'(\d+(?:,\d{3})*(?:\.\d{1,2})?)\s*$', file_name)
        if name_match:
            result["amount"] = float(name_match.group(1).replace(',', ''))
            if result["parseStatus"] == "未提取":
                result["parseStatus"] = "从文件名提取"
    
    if not result["itemName"]:
        name_part = re.sub(r'[\d,_\-\.]+$', '', file_name)
        name_part = re.sub(r'^[_\d]+', '', name_part)
        result["itemName"] = name_part.strip() or file_name.replace('.pdf', '')
    
    result["category"] = categorize(result["itemName"], file_name, text)
    
    return result


def scan_pdfs(directory):
    """递归扫描所有PDF文件（排除输出目录）"""
    output_dir_name = "分类合并结果"
    pdf_files = []
    for root, dirs, files in os.walk(directory):
        # 排除输出目录
        if output_dir_name in root:
            continue
        for f in files:
            if f.lower().endswith('.pdf'):
                pdf_files.append(os.path.join(root, f))
    return pdf_files


def create_excel(invoices, output_file):
    """生成Excel报表"""
    wb = Workbook()
    
    # Sheet 1: 发票明细
    ws1 = wb.active
    ws1.title = "发票明细"
    
    headers = ["序号", "文件路径", "物品/项目", "类别", "金额(元)", "发票号码", "开票日期", "销售方", "解析状态"]
    for col, header in enumerate(headers, 1):
        cell = ws1.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    for row_idx, inv in enumerate(invoices, 2):
        ws1.cell(row=row_idx, column=1, value=row_idx - 1)
        ws1.cell(row=row_idx, column=2, value=inv["filePath"])
        ws1.cell(row=row_idx, column=3, value=inv["itemName"])
        ws1.cell(row=row_idx, column=4, value=inv["category"])
        ws1.cell(row=row_idx, column=5, value=inv["amount"])
        ws1.cell(row=row_idx, column=5).number_format = '#,##0.00'
        ws1.cell(row=row_idx, column=6, value=inv["invoiceNumber"])
        ws1.cell(row=row_idx, column=7, value=inv["date"])
        ws1.cell(row=row_idx, column=8, value=inv["seller"])
        ws1.cell(row=row_idx, column=9, value=inv["parseStatus"])
    
    # Sheet 2: 分类汇总
    ws2 = wb.create_sheet("分类汇总")
    ws2.sheet_properties.tabColor = "52C41A"
    
    headers2 = ["类别", "笔数", "金额(元)", "占比(%)", "平均金额"]
    for col, header in enumerate(headers2, 1):
        cell = ws2.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    
    cat_stats = {}
    for inv in invoices:
        cat = inv["category"]
        if cat not in cat_stats:
            cat_stats[cat] = {"count": 0, "total": 0}
        cat_stats[cat]["count"] += 1
        cat_stats[cat]["total"] += inv["amount"]
    
    total_amount = sum(inv["amount"] for inv in invoices)
    sorted_cats = sorted(cat_stats.items(), key=lambda x: x[1]["total"], reverse=True)
    
    for row_idx, (cat, stats) in enumerate(sorted_cats, 2):
        ws2.cell(row=row_idx, column=1, value=cat)
        ws2.cell(row=row_idx, column=2, value=stats["count"])
        ws2.cell(row=row_idx, column=3, value=stats["total"])
        ws2.cell(row=row_idx, column=3).number_format = '#,##0.00'
        ws2.cell(row=row_idx, column=4, value=stats["total"] / total_amount if total_amount > 0 else 0)
        ws2.cell(row=row_idx, column=4).number_format = '0.0%'
        ws2.cell(row=row_idx, column=5, value=stats["total"] / stats["count"] if stats["count"] > 0 else 0)
        ws2.cell(row=row_idx, column=5).number_format = '#,##0.00'
    
    total_row2 = len(sorted_cats) + 2
    ws2.cell(row=total_row2, column=1, value="合计")
    ws2.cell(row=total_row2, column=2, value=len(invoices))
    ws2.cell(row=total_row2, column=3, value=total_amount)
    ws2.cell(row=total_row2, column=3).number_format = '#,##0.00'
    ws2.cell(row=total_row2, column=4, value=1)
    ws2.cell(row=total_row2, column=4).number_format = '0.0%'
    
    ws1.column_dimensions['A'].width = 8
    ws1.column_dimensions['B'].width = 60
    ws1.column_dimensions['C'].width = 25
    ws1.column_dimensions['D'].width = 14
    ws1.column_dimensions['E'].width = 14
    ws2.column_dimensions['A'].width = 16
    ws2.column_dimensions['B'].width = 10
    ws2.column_dimensions['C'].width = 16
    ws2.column_dimensions['D'].width = 12
    ws2.column_dimensions['E'].width = 16
    
    wb.save(output_file)


def create_html_by_category(invoices, output_dir):
    """按分类生成HTML"""
    from datetime import datetime
    classified = {}
    for inv in invoices:
        cat = inv["category"]
        if cat not in classified:
            classified[cat] = []
        classified[cat].append(inv)
    
    for cat, cat_invoices in classified.items():
        html_content = f"""<!DOCTYPE html>
<html lang="zh-CN">
<head>
    <meta charset="UTF-8">
    <title>{cat} - 发票汇总</title>
    <style>
        * {{ margin: 0; padding: 0; box-sizing: border-box; }}
        body {{ font-family: "Microsoft YaHei", sans-serif; padding: 20px; background: #f5f5f5; color: #333; }}
        .container {{ max-width: 900px; margin: 0 auto; background: #fff; padding: 30px; border-radius: 8px; box-shadow: 0 2px 8px rgba(0,0,0,0.1); }}
        h1 {{ font-size: 22px; color: #2F5496; margin-bottom: 8px; }}
        .meta {{ font-size: 13px; color: #888; margin-bottom: 20px; }}
        .invoice {{ border: 1px solid #e0e0e0; border-radius: 6px; padding: 15px; margin: 10px 0; }}
        .invoice h3 {{ font-size: 14px; color: #2F5496; margin-bottom: 8px; }}
        .invoice p {{ font-size: 13px; margin: 3px 0; }}
        .amount {{ font-size: 16px; font-weight: bold; color: #d46b08; }}
        .total {{ font-size: 18px; font-weight: bold; color: #2F5496; text-align: right; margin-top: 20px; padding-top: 15px; border-top: 2px solid #2F5496; }}
        @media print {{ body {{ background: #fff; padding: 0; }} .container {{ box-shadow: none; }} }}
    </style>
</head>
<body>
<div class="container">
    <h1>{cat} - 发票汇总（共 {len(cat_invoices)} 张）</h1>
    <div class="meta">生成日期: {datetime.now().strftime('%Y-%m-%d %H:%M')}</div>
"""
        for i in range(0, len(cat_invoices), 2):
            inv1 = cat_invoices[i]
            inv2 = cat_invoices[i + 1] if i + 1 < len(cat_invoices) else None
            
            html_content += '<div style="display:flex;gap:15px;margin-bottom:15px;">'
            
            html_content += f'<div class="invoice" style="flex:1;"><h3>发票 {i+1}</h3>'
            html_content += f'<p><b>物品:</b> {inv1["itemName"]}</p>'
            html_content += f'<p class="amount">金额: ¥{inv1["amount"]:.2f}</p>'
            if inv1["invoiceNumber"]:
                html_content += f'<p><b>号码:</b> {inv1["invoiceNumber"]}</p>'
            if inv1["date"]:
                html_content += f'<p><b>日期:</b> {inv1["date"]}</p>'
            if inv1["seller"]:
                html_content += f'<p><b>销售方:</b> {inv1["seller"]}</p>'
            html_content += '</div>'
            
            if inv2:
                html_content += f'<div class="invoice" style="flex:1;"><h3>发票 {i+2}</h3>'
                html_content += f'<p><b>物品:</b> {inv2["itemName"]}</p>'
                html_content += f'<p class="amount">金额: ¥{inv2["amount"]:.2f}</p>'
                if inv2["invoiceNumber"]:
                    html_content += f'<p><b>号码:</b> {inv2["invoiceNumber"]}</p>'
                if inv2["date"]:
                    html_content += f'<p><b>日期:</b> {inv2["date"]}</p>'
                if inv2["seller"]:
                    html_content += f'<p><b>销售方:</b> {inv2["seller"]}</p>'
                html_content += '</div>'
            
            html_content += '</div>'
        
        cat_total = sum(inv["amount"] for inv in cat_invoices)
        html_content += f'<div class="total">总金额: ¥{cat_total:.2f}</div>'
        html_content += '</div></body></html>'
        
        output_file = os.path.join(output_dir, f"{cat}.html")
        with open(output_file, 'w', encoding='utf-8') as f:
            f.write(html_content)


def main():
    if len(sys.argv) < 2:
        print("用法: python 发票统计.py <文件夹路径>")
        sys.exit(1)
    
    folder = sys.argv[1]
    if not os.path.isdir(folder):
        print(f"错误: 文件夹不存在 - {folder}")
        sys.exit(1)
    
    print(f"\n{'=' * 60}")
    print(f"  发票统计 + 分类合并工具")
    print(f"{'=' * 60}\n")
    print(f"扫描文件夹: {folder}\n")
    
    # 扫描PDF
    pdf_files = scan_pdfs(folder)
    print(f"找到 {len(pdf_files)} 个PDF文件（含所有子文件夹）\n")
    
    if not pdf_files:
        print("未找到任何PDF文件。")
        return
    
    # 解析
    print("正在解析发票...")
    invoices = []
    for i, pdf_path in enumerate(pdf_files, 1):
        inv = parse_invoice(pdf_path)
        invoices.append(inv)
        status = f"¥{inv['amount']:.2f}" if inv["amount"] > 0 else inv["parseStatus"]
        print(f"[{i}/{len(pdf_files)}] {inv['fileName']} -> {inv['category']} | {status}")
    
    # 输出目录
    output_dir = os.path.join(folder, "分类合并结果")
    os.makedirs(output_dir, exist_ok=True)
    
    # 保存JSON数据
    json_file = os.path.join(output_dir, "invoices.json")
    with open(json_file, 'w', encoding='utf-8') as f:
        json.dump(invoices, f, ensure_ascii=False, indent=2)
    
    # 生成Excel
    from datetime import datetime
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    excel_file = os.path.join(output_dir, f"发票分类统计_{timestamp}.xlsx")
    print(f"\n生成Excel: {excel_file}")
    create_excel(invoices, excel_file)
    
    # 生成HTML
    print("生成分类HTML...")
    create_html_by_category(invoices, output_dir)
    
    # 合并PDF
    print("合并分类PDF（A4每页两张发票）...")
    import subprocess
    merge_script = os.path.join(os.path.dirname(os.path.abspath(__file__)), "merge_pdfs.py")
    if os.path.exists(merge_script):
        subprocess.run([sys.executable, merge_script, os.path.join(output_dir, "invoices.json"), output_dir])
    
    # 统计
    total_amount = sum(inv["amount"] for inv in invoices)
    valid_count = sum(1 for inv in invoices if inv["amount"] > 0)
    
    print(f"\n{'=' * 60}")
    print(f"  完成!")
    print(f"{'=' * 60}")
    print(f"  总发票数: {len(invoices)} 张")
    print(f"  成功解析: {valid_count} 张")
    print(f"  总金额: ¥{total_amount:.2f}")
    print(f"  输出目录: {output_dir}")
    print(f"{'=' * 60}\n")


if __name__ == "__main__":
    main()
