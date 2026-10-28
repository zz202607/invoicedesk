"""Render the invoice pool export as an Excel workbook.

Sheets: invoice detail, category summary, company (buyer) summary, and export
metadata. The workbook is meant to be submitted, so it never contains absolute
source paths — only the source file name and page number.
"""

from __future__ import annotations

from collections.abc import Callable, Sequence
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from app.core.export import ExportError, ExportFilters, ExportRow, ExportSummary
from app.core.print_pdf import PrintInfo
from app.core.text import normalize

EXCEL_FILENAME = "发票明细.xlsx"

DETAIL_HEADERS = (
    "序号",
    "开票日期",
    "发票号码",
    "票型",
    "销售方",
    "购买方",
    "项目摘要",
    "分类",
    "不含税金额",
    "税额",
    "价税合计",
    "待确认",
    "低置信字段",
    "备注",
    "来源文件",
    "页码",
)
DETAIL_WIDTHS = (6, 12, 24, 16, 26, 22, 30, 12, 14, 14, 14, 8, 22, 30, 28, 8)
MONEY_FORMAT = "#,##0.00"
TEXT_FORMAT = "@"
MAX_CELL_CHARS = 500

HEADER_FILL = PatternFill(fill_type="solid", fgColor="1F4E79")
HEADER_FONT = Font(name="微软雅黑", bold=True, color="FFFFFF")
BODY_FONT = Font(name="微软雅黑")
BOLD_FONT = Font(name="微软雅黑", bold=True)
REVIEW_FILL = PatternFill(fill_type="solid", fgColor="FFF2CC")
REVIEW_FONT = Font(name="微软雅黑", bold=True, color="C00000")
_THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)
TOP_BORDER = Border(top=Side(style="medium", color="7F7F7F"))
MONEY_ALIGN = Alignment(horizontal="right", vertical="top")
CENTER_ALIGN = Alignment(horizontal="center", vertical="top")
TEXT_ALIGN = Alignment(horizontal="left", vertical="top", wrap_text=True)


def _text(value: str, limit: int = MAX_CELL_CHARS) -> str:
    text = normalize(value)
    if len(text) > limit:
        return text[: limit - 1] + "…"
    return text


def _set_text(cell, value: str, *, limit: int = MAX_CELL_CHARS) -> None:
    text = _text(value, limit)
    cell.value = text
    # Invoice text (now editable in the workbench) is untrusted; force the
    # cell to text so leading =, +, -, @, tab, or CR can never be interpreted
    # as a formula when the workbook is opened.
    cell.data_type = "s"


def _set_literal(cell, value: str, *, limit: int = 1000) -> None:
    """Write application-authored text without NFKC folding its punctuation."""
    text = value.replace("\r", " ").replace("\n", " ").strip()
    if len(text) > limit:
        text = text[: limit - 1] + "…"
    cell.value = text
    cell.data_type = "s"


def _set_money(cell, cents: int | None) -> None:
    cell.value = None if cents is None else round(cents / 100, 2)
    cell.number_format = MONEY_FORMAT
    cell.alignment = MONEY_ALIGN


def _set_count(cell, value: int) -> None:
    cell.value = value
    cell.number_format = "0"
    cell.alignment = CENTER_ALIGN


def _write_header(worksheet: Worksheet, headers: Sequence[str]) -> None:
    for column, title in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=column, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    worksheet.row_dimensions[1].height = 24


def _set_widths(worksheet: Worksheet, widths: Sequence[int]) -> None:
    for column, width in enumerate(widths, start=1):
        worksheet.column_dimensions[get_column_letter(column)].width = width


def _finish_sheet(worksheet: Worksheet, widths: Sequence[int], *, landscape: bool = False) -> None:
    _set_widths(worksheet, widths)
    worksheet.sheet_view.showGridLines = False
    worksheet.freeze_panes = "A2"
    if landscape:
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.paperSize = Worksheet.PAPERSIZE_A4
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.print_title_rows = "1:1"


def _write_detail(worksheet: Worksheet, rows: Sequence[ExportRow], summary: ExportSummary) -> None:
    _write_header(worksheet, DETAIL_HEADERS)
    for index, row in enumerate(rows, start=1):
        excel_row = index + 1
        sequence = worksheet.cell(row=excel_row, column=1)
        _set_count(sequence, index)

        date_cell = worksheet.cell(row=excel_row, column=2)
        _set_text(date_cell, row.invoice_date)
        date_cell.number_format = TEXT_FORMAT
        date_cell.alignment = CENTER_ALIGN

        number_cell = worksheet.cell(row=excel_row, column=3)
        _set_text(number_cell, row.invoice_number)
        number_cell.number_format = TEXT_FORMAT

        type_cell = worksheet.cell(row=excel_row, column=4)
        _set_text(type_cell, row.type_label)
        type_cell.alignment = CENTER_ALIGN

        seller_cell = worksheet.cell(row=excel_row, column=5)
        _set_text(seller_cell, row.seller_name)
        buyer_cell = worksheet.cell(row=excel_row, column=6)
        _set_text(buyer_cell, row.buyer_name)
        item_cell = worksheet.cell(row=excel_row, column=7)
        _set_text(item_cell, row.item_summary)
        category_cell = worksheet.cell(row=excel_row, column=8)
        _set_text(category_cell, row.category)
        for cell in (seller_cell, buyer_cell, item_cell, category_cell):
            cell.alignment = TEXT_ALIGN

        _set_money(worksheet.cell(row=excel_row, column=9), row.amount_net_cents)
        _set_money(worksheet.cell(row=excel_row, column=10), row.tax_amount_cents)
        _set_money(worksheet.cell(row=excel_row, column=11), row.amount_total_cents)

        review_cell = worksheet.cell(row=excel_row, column=12)
        review_cell.value = "是" if row.needs_review else ""
        review_cell.alignment = CENTER_ALIGN
        if row.needs_review:
            review_cell.fill = REVIEW_FILL
            review_cell.font = REVIEW_FONT

        confidence_cell = worksheet.cell(row=excel_row, column=13)
        _set_literal(confidence_cell, row.confidence_notes)
        confidence_cell.alignment = TEXT_ALIGN

        warning_cell = worksheet.cell(row=excel_row, column=14)
        _set_text(warning_cell, row.warning_text)
        warning_cell.alignment = TEXT_ALIGN
        source_cell = worksheet.cell(row=excel_row, column=15)
        _set_text(source_cell, row.source_name)
        source_cell.alignment = TEXT_ALIGN
        page_cell = worksheet.cell(row=excel_row, column=16)
        _set_text(page_cell, row.page_label)
        page_cell.alignment = CENTER_ALIGN

        for column in range(1, len(DETAIL_HEADERS) + 1):
            cell = worksheet.cell(row=excel_row, column=column)
            cell.border = BORDER
            cell.font = BODY_FONT
        if row.needs_review:
            review_cell.font = REVIEW_FONT

    total_row = len(rows) + 3
    label = worksheet.cell(row=total_row, column=8)
    label.value = f"合计（{summary.count} 张）"
    label.font = BOLD_FONT
    label.alignment = MONEY_ALIGN
    for column, cents in ((9, summary.net_cents), (10, summary.tax_cents), (11, summary.total_cents)):
        cell = worksheet.cell(row=total_row, column=column)
        _set_money(cell, cents)
        cell.font = BOLD_FONT
        cell.border = TOP_BORDER
    label.border = TOP_BORDER

    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(DETAIL_HEADERS))}{len(rows) + 1}"
    _finish_sheet(worksheet, DETAIL_WIDTHS, landscape=True)


def _group(rows: Sequence[ExportRow], key: Callable[[ExportRow], str], fallback: str) -> dict:
    groups: dict[str, dict[str, int]] = {}
    for row in rows:
        name = key(row) or fallback
        aggregate = groups.setdefault(
            name, {"count": 0, "review": 0, "net": 0, "tax": 0, "total": 0}
        )
        aggregate["count"] += 1
        aggregate["review"] += int(row.needs_review)
        aggregate["net"] += row.amount_net_cents or 0
        aggregate["tax"] += row.tax_amount_cents or 0
        aggregate["total"] += row.amount_total_cents or 0
    return groups


def _write_category_summary(
    worksheet: Worksheet, rows: Sequence[ExportRow], summary: ExportSummary
) -> None:
    headers = ("分类", "张数", "不含税金额", "税额", "价税合计", "占比")
    _write_header(worksheet, headers)
    groups = _group(rows, lambda row: row.category, "未分类")
    ordered = sorted(groups.items(), key=lambda item: (-item[1]["total"], item[0]))

    for index, (name, aggregate) in enumerate(ordered, start=2):
        cell = worksheet.cell(row=index, column=1)
        _set_text(cell, name)
        _set_count(worksheet.cell(row=index, column=2), aggregate["count"])
        _set_money(worksheet.cell(row=index, column=3), aggregate["net"])
        _set_money(worksheet.cell(row=index, column=4), aggregate["tax"])
        _set_money(worksheet.cell(row=index, column=5), aggregate["total"])
        share = worksheet.cell(row=index, column=6)
        share.value = (
            aggregate["total"] / summary.total_cents if summary.total_cents else 0.0
        )
        share.number_format = "0.0%"
        share.alignment = MONEY_ALIGN
        for column in range(1, 7):
            cell = worksheet.cell(row=index, column=column)
            cell.border = BORDER
            cell.font = BODY_FONT

    total_row = len(groups) + 2
    label = worksheet.cell(row=total_row, column=1, value=f"合计（{summary.count} 张）")
    label.font = BOLD_FONT
    label.border = TOP_BORDER
    count_cell = worksheet.cell(row=total_row, column=2)
    _set_count(count_cell, summary.count)
    count_cell.font = BOLD_FONT
    count_cell.border = TOP_BORDER
    for column, cents in ((3, summary.net_cents), (4, summary.tax_cents), (5, summary.total_cents)):
        cell = worksheet.cell(row=total_row, column=column)
        _set_money(cell, cents)
        cell.font = BOLD_FONT
        cell.border = TOP_BORDER
    share = worksheet.cell(row=total_row, column=6)
    share.value = 1.0 if summary.total_cents else 0.0
    share.number_format = "0.0%"
    share.font = BOLD_FONT
    share.alignment = MONEY_ALIGN
    share.border = TOP_BORDER

    worksheet.auto_filter.ref = f"A1:F{len(groups) + 1}"
    _finish_sheet(worksheet, (18, 10, 15, 14, 15, 10))


def _write_buyer_summary(
    worksheet: Worksheet, rows: Sequence[ExportRow], summary: ExportSummary
) -> None:
    headers = ("购买方", "张数", "待确认张数", "价税合计")
    _write_header(worksheet, headers)
    groups = _group(rows, lambda row: row.buyer_name, "未识别")
    ordered = sorted(groups.items(), key=lambda item: (-item[1]["total"], item[0]))

    for index, (name, aggregate) in enumerate(ordered, start=2):
        cell = worksheet.cell(row=index, column=1)
        _set_text(cell, name)
        _set_count(worksheet.cell(row=index, column=2), aggregate["count"])
        _set_count(worksheet.cell(row=index, column=3), aggregate["review"])
        _set_money(worksheet.cell(row=index, column=4), aggregate["total"])
        for column in range(1, 5):
            cell = worksheet.cell(row=index, column=column)
            cell.border = BORDER
            cell.font = BODY_FONT

    total_row = len(groups) + 2
    label = worksheet.cell(row=total_row, column=1, value=f"合计（{summary.count} 张）")
    label.font = BOLD_FONT
    label.border = TOP_BORDER
    count_cell = worksheet.cell(row=total_row, column=2)
    _set_count(count_cell, summary.count)
    count_cell.font = BOLD_FONT
    count_cell.border = TOP_BORDER
    review_cell = worksheet.cell(row=total_row, column=3)
    _set_count(review_cell, summary.review_count)
    review_cell.font = BOLD_FONT
    review_cell.border = TOP_BORDER
    total_cell = worksheet.cell(row=total_row, column=4)
    _set_money(total_cell, summary.total_cents)
    total_cell.font = BOLD_FONT
    total_cell.border = TOP_BORDER

    worksheet.auto_filter.ref = f"A1:D{len(groups) + 1}"
    _finish_sheet(worksheet, (30, 10, 12, 15))


def _write_metadata(
    worksheet: Worksheet,
    *,
    filters: ExportFilters,
    summary: ExportSummary,
    generated_at: str,
    print_info: PrintInfo | None,
) -> None:
    worksheet.merge_cells("A1:B1")
    title = worksheet.cell(row=1, column=1, value="发票导出说明")
    title.font = Font(name="微软雅黑", size=14, bold=True)
    title.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 26

    entries: list[tuple[str, object, str | None]] = [
        ("导出时间", generated_at, None),
        ("筛选条件", filters.describe(), None),
        ("待确认处理", filters.review_mode(), None),
        ("发票条数", summary.count, "0"),
        ("待确认条数", summary.review_count, "0"),
        ("金额缺失条数", summary.missing_amount_count, "0"),
        ("不含税合计", round(summary.net_cents / 100, 2), MONEY_FORMAT),
        ("税额合计", round(summary.tax_cents / 100, 2), MONEY_FORMAT),
        ("价税合计", round(summary.total_cents / 100, 2), MONEY_FORMAT),
        ("来源票面数", summary.unique_pages, "0"),
    ]
    if print_info is not None:
        entries.extend(
            [
                ("二合一打印页数", print_info.pages, "0"),
                ("已打印票面数", print_info.unique_pages, "0"),
                ("缺源文件票面数", len(print_info.skipped_missing), "0"),
                ("无法打开票面数", len(print_info.skipped_invalid), "0"),
            ]
        )

    for index, (key, value, number_format) in enumerate(entries, start=2):
        key_cell = worksheet.cell(row=index, column=1, value=key)
        key_cell.font = BOLD_FONT
        key_cell.alignment = Alignment(vertical="top")
        value_cell = worksheet.cell(row=index, column=2)
        if isinstance(value, str):
            _set_literal(value_cell, value)
        else:
            value_cell.value = value
        if number_format:
            value_cell.number_format = number_format
            value_cell.alignment = MONEY_ALIGN
        else:
            value_cell.alignment = TEXT_ALIGN

    note_row = len(entries) + 3
    notes = (
        "说明：",
        "来源文件列仅保留文件名和页码，不包含本机路径。",
        "待确认发票请在提交前人工核对；“备注”列列出识别警告。",
        "本表由 InvoiceDesk 在本机离线生成，未上传任何数据。",
    )
    for offset, note in enumerate(notes):
        cell = worksheet.cell(row=note_row + offset, column=1)
        _set_literal(cell, note, limit=200)
        cell.font = BOLD_FONT if offset == 0 else BODY_FONT

    _set_widths(worksheet, (22, 70))
    worksheet.sheet_view.showGridLines = False


def write_workbook(
    path: str | Path,
    rows: Sequence[ExportRow],
    summary: ExportSummary,
    *,
    filters: ExportFilters,
    generated_at: str,
    print_info: PrintInfo | None = None,
) -> None:
    """Write the Excel export package. Raises ``OSError`` on write failures."""
    workbook = Workbook()
    workbook.properties.creator = "InvoiceDesk"
    workbook.properties.lastModifiedBy = "InvoiceDesk"
    workbook.properties.title = "发票明细"
    workbook.properties.description = "本机离线导出的发票数据，含财务信息。"

    detail = workbook.active
    detail.title = "发票明细"
    _write_detail(detail, rows, summary)

    _write_category_summary(workbook.create_sheet("分类汇总"), rows, summary)
    _write_buyer_summary(workbook.create_sheet("公司汇总"), rows, summary)
    _write_metadata(
        workbook.create_sheet("导出信息"),
        filters=filters,
        summary=summary,
        generated_at=generated_at,
        print_info=print_info,
    )

    target = Path(path)
    if target.is_symlink():
        raise ExportError("导出目标是符号链接，已拒绝写入。")
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
