"""CLI tests for ``report``: build a sheet, edit lines, and export."""

from __future__ import annotations

import sqlite3
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import load_workbook

from app.cli import main
from tests.pdf_factory import full_digital_lines, make_pdf


def invoice_lines(
    *,
    number: str = "26502000000000000001",
    buyer: str = "示例甲公司",
    date: str = "2026年07月09日",
    seller: str = "示例销售方",
    item: str = "*生产生活服务*办公用品",
    small: str = "¥ 120.00",
    uppercase: str = "壹佰贰拾圆整",
    net: str = "¥113.21",
    tax: str = "¥6.79",
) -> list[tuple[float, float, str]]:
    lines = full_digital_lines(
        number=number,
        seller=seller,
        item=item,
        small=small,
        uppercase=uppercase,
        net=net,
        tax=tax,
    )
    replacements = {"示例甲公司": buyer, "2026年07月30日": date}
    return [(x, y, replacements.get(text, text)) for x, y, text in lines]


def write_invoice(folder: Path, name: str, **kwargs) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    return make_pdf(folder, invoice_lines(**kwargs), name=name)


def prepare_pool(tmp_path: Path, count: int = 1) -> Path:
    folder = tmp_path / "july"
    for index in range(1, count + 1):
        write_invoice(
            folder,
            f"invoice{index}.pdf",
            number=f"2650200000000000000{index}",
            date=f"2026年07月{8 + index:02d}日",
            seller=f"示例销售方{index}",
        )
    database = tmp_path / "pool.db"
    assert main(["import", str(folder), "--db", str(database)]) == 0
    return database


def money(cell) -> Decimal:
    return Decimal(str(cell.value)).quantize(Decimal("0.01"))


def create_report(database: Path, **overrides) -> int:
    args = [
        "report",
        "create",
        "--db",
        str(database),
        "--company",
        "示例公司",
        "--reporter",
        "张三",
        "--month",
        "2026-07",
        "--department",
        "示例部门",
    ]
    for key, value in overrides.items():
        args.extend([f"--{key.replace('_', '-')}", str(value)])
    return main(args)


def test_report_cli_builds_and_exports_a_sheet(tmp_path: Path, capsys) -> None:
    database = prepare_pool(tmp_path)
    capsys.readouterr()

    assert create_report(database) == 0
    captured = capsys.readouterr()
    assert "已创建报销单 #1" in captured.out

    assert (
        main(
            [
                "report",
                "from-invoices",
                "--db",
                str(database),
                "--report",
                "1",
                "--project",
                "示例项目",
            ]
        )
        == 0
    )
    capsys.readouterr()

    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-19",
                "--category",
                "转账",
                "--amount",
                "-1000.00",
                "--description",
                "示例转账",
            ]
        )
        == 0
    )
    assert (
        main(
            [
                "report",
                "add-mileage",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-10",
                "--km",
                "150",
                "--description",
                "示例行程",
                "--project",
                "示例项目二",
            ]
        )
        == 0
    )
    capsys.readouterr()

    output = tmp_path / "report_out"
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(output),
                "--layout",
                "company",
            ]
        )
        == 0
    )
    captured = capsys.readouterr()
    assert "已导出报销单 #1：3 行" in captured.out

    target = output / "示例公司费用报销支出明细表（张三）2607.xlsx"
    assert target.exists()
    worksheet = load_workbook(target).active
    assert worksheet["A1"].value == "示例公司7月报销费用支出明细表"
    assert money(worksheet["B3"]) == Decimal("-655.00")
    assert [worksheet.cell(row=5, column=column).value for column in (1, 2, 6)] == [
        1,
        "2026.7.9",
        "有票",
    ]
    assert money(worksheet["D5"]) == Decimal("120.00")
    assert worksheet["H5"].value == "示例公司"
    assert "办公用品" in worksheet["G5"].value
    assert [worksheet.cell(row=6, column=column).value for column in (2, 3, 6)] == [
        "2026.7.19",
        "转账",
        "无票",
    ]
    assert money(worksheet["D6"]) == Decimal("-1000.00")
    assert [worksheet.cell(row=7, column=column).value for column in (2, 3, 6)] == [
        "2026.7.10",
        "汽车费用",
        "无票",
    ]
    assert money(worksheet["D7"]) == Decimal("225.00")
    assert worksheet.max_row == 7


def test_report_cli_list_show_update_and_remove(tmp_path: Path, capsys) -> None:
    database = tmp_path / "pool.db"
    assert create_report(database) == 0
    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--category",
                "办公费用",
                "--amount",
                "120",
                "--description",
                "办公用品",
            ]
        )
        == 0
    )
    assert (
        main(
            [
                "report",
                "add-mileage",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-10",
                "--km",
                "150",
                "--description",
                "示例行程",
            ]
        )
        == 0
    )
    capsys.readouterr()

    assert main(["report", "list", "--db", str(database)]) == 0
    assert "2 行" in capsys.readouterr().out

    assert main(["report", "show", "--db", str(database), "--report", "1"]) == 0
    captured = capsys.readouterr()
    assert "共 2 行，合计 345.00 元" in captured.out
    assert "[150km] 示例行程" in captured.out
    assert "2026.7.9" in captured.out

    assert (
        main(
            [
                "report",
                "update-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "1",
                "--amount",
                "200",
                "--description",
                "修正后的说明",
            ]
        )
        == 0
    )
    assert main(["report", "show", "--db", str(database), "--report", "1"]) == 0
    captured = capsys.readouterr()
    assert "修正后的说明" in captured.out
    assert "共 2 行，合计 425.00 元" in captured.out

    assert (
        main(
            [
                "report",
                "remove-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "2",
            ]
        )
        == 0
    )
    assert main(["report", "show", "--db", str(database), "--report", "1"]) == 0
    assert "共 1 行，合计 200.00 元" in capsys.readouterr().out

    assert (
        main(
            [
                "report",
                "update-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "1",
            ]
        )
        == 2
    )
    assert "没有要修改" in capsys.readouterr().err


def test_report_cli_links_and_unlinks_invoices(tmp_path: Path, capsys) -> None:
    database = prepare_pool(tmp_path)
    capsys.readouterr()
    assert create_report(database) == 0
    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "120",
                "--description",
                "办公用品",
            ]
        )
        == 0
    )
    capsys.readouterr()

    assert (
        main(
            [
                "report",
                "link",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "1",
                "--invoice",
                "1",
            ]
        )
        == 0
    )
    assert "已关联 1 张发票" in capsys.readouterr().out
    assert main(["report", "show", "--db", str(database), "--report", "1"]) == 0
    assert "发票 #1" in capsys.readouterr().out

    assert (
        main(
            [
                "report",
                "link",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "1",
                "--invoice",
                "1",
            ]
        )
        == 0
    )
    assert "没有新增关联" in capsys.readouterr().out

    assert (
        main(
            [
                "report",
                "unlink",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "1",
                "--invoice",
                "1",
            ]
        )
        == 0
    )
    assert "已取消 1 张发票关联" in capsys.readouterr().out

    assert (
        main(
            [
                "report",
                "link",
                "--db",
                str(database),
                "--report",
                "1",
                "--line",
                "1",
                "--invoice",
                "999",
            ]
        )
        == 2
    )
    assert "发票不存在" in capsys.readouterr().err


def test_report_cli_from_invoices_is_idempotent(tmp_path: Path, capsys) -> None:
    database = prepare_pool(tmp_path, count=2)
    capsys.readouterr()
    assert create_report(database) == 0

    assert (
        main(
            [
                "report",
                "from-invoices",
                "--db",
                str(database),
                "--report",
                "1",
            ]
        )
        == 0
    )
    assert "已从发票池添加 2 行" in capsys.readouterr().out

    assert (
        main(
            [
                "report",
                "from-invoices",
                "--db",
                str(database),
                "--report",
                "1",
            ]
        )
        == 1
    )
    captured = capsys.readouterr()
    assert "没有可添加的发票" in captured.out
    assert "跳过 2 张：已关联 2" in captured.err


def test_report_cli_export_project_header(tmp_path: Path, capsys) -> None:
    database = tmp_path / "pool.db"
    assert create_report(database) == 0
    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "120",
                "--description",
                "办公用品",
                "--project",
                "示例项目",
            ]
        )
        == 0
    )
    capsys.readouterr()

    output = tmp_path / "out"
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(output),
                "--project-header",
                "费用发生项目",
            ]
        )
        == 0
    )
    worksheet = load_workbook(next(output.glob("*.xlsx"))).active
    assert worksheet.cell(row=4, column=7).value == "费用发生项目"
    assert worksheet["G5"].value == "示例项目"


def test_report_cli_export_warns_on_mismatch_and_unlinked(tmp_path: Path, capsys) -> None:
    database = prepare_pool(tmp_path)
    capsys.readouterr()
    assert create_report(database) == 0

    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "100",
                "--description",
                "示例费用",
                "--invoice",
                "1",
            ]
        )
        == 0
    )
    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "50",
                "--description",
                "示例费用二",
                "--has-invoice",
            ]
        )
        == 0
    )
    capsys.readouterr()

    output = tmp_path / "out"
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(output),
            ]
        )
        == 0
    )
    captured = capsys.readouterr()
    assert "关联发票合计与行金额不一致" in captured.err
    assert "标记有票但未关联发票" in captured.err

    worksheet = load_workbook(next(output.glob("*.xlsx"))).active
    assert worksheet["D5"].font.bold is True
    assert worksheet["D5"].font.color.rgb.endswith("C00000")
    assert worksheet["D6"].font.bold is False


def test_report_cli_export_guards(tmp_path: Path, capsys, monkeypatch) -> None:
    database = tmp_path / "pool.db"
    assert create_report(database) == 0
    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "120",
                "--description",
                "办公用品",
            ]
        )
        == 0
    )
    capsys.readouterr()

    fake_root = tmp_path / "repo"
    fake_root.mkdir()
    monkeypatch.setattr("app.core.paths.repo_root", lambda: fake_root)
    unsafe = fake_root / "reports"
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(unsafe),
            ]
        )
        == 2
    )
    assert "仓库内非忽略目录" in capsys.readouterr().err

    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(unsafe),
                "--force",
            ]
        )
        == 0
    )
    assert any(unsafe.glob("*.xlsx"))

    monkeypatch.setattr("app.cli_report.remote_storage_reason", lambda path: "OneDrive")
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(tmp_path / "remote"),
            ]
        )
        == 2
    )
    assert "位于OneDrive" in capsys.readouterr().err
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(tmp_path / "remote"),
                "--force",
            ]
        )
        == 0
    )


def test_report_cli_export_errors(tmp_path: Path, capsys, monkeypatch) -> None:
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(tmp_path / "missing.db"),
                "--report",
                "1",
                "--out",
                str(tmp_path / "out"),
            ]
        )
        == 1
    )
    assert "数据库不存在" in capsys.readouterr().err

    database = tmp_path / "pool.db"
    assert create_report(database) == 0
    capsys.readouterr()
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "9",
                "--out",
                str(tmp_path / "out"),
            ]
        )
        == 2
    )
    assert "报销单不存在" in capsys.readouterr().err

    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(tmp_path / "out"),
            ]
        )
        == 1
    )
    assert "没有报销行" in capsys.readouterr().err

    assert (
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "120",
                "--description",
                "办公用品",
            ]
        )
        == 0
    )
    capsys.readouterr()

    target = tmp_path / "out" / "示例公司费用报销支出明细表（张三）2607.xlsx"
    original = Path.is_symlink

    def fake_is_symlink(self):
        return self == target or original(self)

    monkeypatch.setattr(Path, "is_symlink", fake_is_symlink)
    assert (
        main(
            [
                "report",
                "export",
                "--db",
                str(database),
                "--report",
                "1",
                "--out",
                str(tmp_path / "out"),
            ]
        )
        == 2
    )
    assert "符号链接" in capsys.readouterr().err
    assert not target.exists()


def test_report_cli_rejects_bad_arguments(tmp_path: Path, capsys) -> None:
    database = tmp_path / "pool.db"
    with pytest.raises(SystemExit) as excinfo:
        main(
            [
                "report",
                "create",
                "--db",
                str(database),
                "--company",
                "示例公司",
                "--reporter",
                "张三",
                "--month",
                "2026-13",
            ]
        )
    assert excinfo.value.code == 2
    capsys.readouterr()

    assert create_report(database) == 0
    capsys.readouterr()

    with pytest.raises(SystemExit):
        main(
            [
                "report",
                "add-line",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-09",
                "--amount",
                "1.234",
                "--description",
                "x",
            ]
        )
    capsys.readouterr()

    with pytest.raises(SystemExit):
        main(
            [
                "report",
                "add-mileage",
                "--db",
                str(database),
                "--report",
                "1",
                "--date",
                "2026-07-10",
                "--km",
                "0",
                "--description",
                "x",
            ]
        )
    capsys.readouterr()

    assert main(["report", "show", "--db", str(database), "--report", "99"]) == 2
    assert "报销单不存在" in capsys.readouterr().err

    assert (
        main(
            [
                "report",
                "from-invoices",
                "--db",
                str(database),
                "--report",
                "99",
            ]
        )
        == 2
    )
    capsys.readouterr()


def test_report_cli_custom_title_and_department(tmp_path: Path, capsys) -> None:
    database = tmp_path / "pool.db"
    assert (
        create_report(
            database,
            company="示例甲",
            reporter="李四",
            month="2026-09",
            department="示例部门乙",
            title="示例甲自定义报销单",
        )
        == 0
    )
    captured = capsys.readouterr()
    assert "示例甲自定义报销单" in captured.out
    assert main(["report", "list", "--db", str(database)]) == 0
    assert "示例甲自定义报销单" in capsys.readouterr().out

    with sqlite3.connect(database) as connection:
        row = connection.execute("SELECT department FROM reports").fetchone()
    assert row == ("示例部门乙",)


def test_report_cli_from_invoices_filters_batch(tmp_path: Path, capsys) -> None:
    database = prepare_pool(tmp_path, count=1)
    capsys.readouterr()
    second = tmp_path / "august"
    write_invoice(
        second,
        "august.pdf",
        number="26502000000000000009",
        date="2026年08月01日",
        seller="示例销售方9",
    )
    assert main(["import", str(second), "--db", str(database)]) == 0
    capsys.readouterr()
    assert create_report(database) == 0
    assert (
        main(
            [
                "report",
                "from-invoices",
                "--db",
                str(database),
                "--report",
                "1",
                "--batch",
                "2",
            ]
        )
        == 0
    )
    assert "已从发票池添加 1 行" in capsys.readouterr().out
    assert main(["report", "show", "--db", str(database), "--report", "1"]) == 0
    assert "2026.8.1" in capsys.readouterr().out
