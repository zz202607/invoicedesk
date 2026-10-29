"""Report Excel rendering: layout presets, values, and write guards."""

from __future__ import annotations

from dataclasses import replace
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from app.core import report_export, reports
from app.core.report_export import (
    display_date,
    report_filename,
    sanitize_filename,
    write_report_workbook,
)
from app.core.reports import ReportError
from tests.report_factory import add_invoice, add_line, make_db, make_report


@pytest.mark.parametrize("value", ["=1+1", "+1", "-1", "@SUM(A1)"])
def test_user_text_is_always_written_as_text(value: str) -> None:
    workbook = Workbook()
    cell = workbook.active["A1"]
    report_export._set_literal(cell, value)
    assert cell.data_type == "s"


@pytest.fixture
def connection(tmp_path):
    conn = make_db(tmp_path)
    yield conn
    conn.close()


def money(cell) -> Decimal:
    return Decimal(str(cell.value)).quantize(Decimal("0.01"))


def build_report(connection) -> tuple[reports.Report, list[reports.ReportLine]]:
    report_id = make_report(connection)
    add_line(
        connection,
        report_id,
        line_date="2026-07-09",
        category="办公费用",
        description="示例费用甲",
        amount_cents=12000,
        project_name="示例项目",
        has_invoice=True,
    )
    add_line(
        connection,
        report_id,
        line_date="2026-07-19",
        category="转账",
        description="示例转账",
        amount_cents=-100000,
    )
    reports.add_mileage(
        connection,
        report_id,
        line_date="2026-07-10",
        km=Decimal(150),
        description="示例行程",
        project_name="示例项目二",
    )
    return reports.get_report(connection, report_id), reports.report_lines(
        connection, report_id
    )


def test_project_layout_matches_sheet_structure(connection, tmp_path) -> None:
    report, lines = build_report(connection)
    target = tmp_path / "report.xlsx"
    write_report_workbook(target, report, lines)

    workbook = load_workbook(target)
    worksheet = workbook.active
    assert worksheet.title == "明细"
    assert worksheet["A1"].value == "示例公司7月报销费用支出明细表"
    assert worksheet["A1"].font.size == 16
    assert worksheet["A1"].font.bold
    assert {str(rng) for rng in worksheet.merged_cells.ranges} == {
        "A1:H1",
        "A2:H2",
        "C3:H3",
        "A8:C8",
        "E8:H8",
    }
    assert worksheet["A2"].value.startswith("部门：示例部门")
    assert "报销人：张三" in worksheet["A2"].value
    assert worksheet["A3"].value == "报销金额："
    assert money(worksheet["B3"]) == Decimal("-655.00")
    assert worksheet["C3"].value == "大写：负陆佰伍拾伍元整"
    assert worksheet["B3"].number_format == "#,##0.00"

    headers = [worksheet.cell(row=4, column=column).value for column in range(1, 9)]
    assert headers == [
        "序号",
        "日期",
        "费用名称",
        "金额",
        "报销人",
        "是否持有发票",
        "项目名称",
        "备注（报销费用说明）",
    ]

    first = [worksheet.cell(row=5, column=column).value for column in range(1, 9)]
    assert first[:3] == [1, "2026.7.9", "办公费用"]
    assert money(worksheet["D5"]) == Decimal("120.00")
    assert first[4:] == ["张三", "有票", "示例项目", "示例费用甲"]
    second = [worksheet.cell(row=6, column=column).value for column in range(1, 9)]
    assert second[:3] == [2, "2026.7.19", "转账"]
    assert money(worksheet["D6"]) == Decimal("-1000.00")
    assert second[4:6] == ["张三", "无票"]
    third = [worksheet.cell(row=7, column=column).value for column in range(1, 9)]
    assert third[:3] == [3, "2026.7.10", "汽车费用"]
    assert money(worksheet["D7"]) == Decimal("225.00")
    assert third[4:] == ["张三", "无票", "示例项目二", "示例行程"]

    assert worksheet.max_row == 8
    assert worksheet["A8"].value == "合计"
    assert money(worksheet["D8"]) == Decimal("-655.00")
    assert worksheet["E8"].value == "大写：负陆佰伍拾伍元整"
    assert worksheet.column_dimensions["H"].width == 42.0
    assert worksheet.sheet_view.showGridLines is False
    assert worksheet.print_title_rows == "$1:$4"

    summary = workbook["分类汇总"]
    assert summary["A1"].value == "示例公司7月报销费用支出明细表 · 分类汇总"
    assert [summary.cell(row=2, column=column).value for column in range(1, 5)] == [
        "费用类别",
        "行数",
        "金额",
        "占比",
    ]
    assert summary["A3"].value == "汽车费用"
    assert summary["B3"].value == 1
    assert money(summary["C3"]) == Decimal("225.00")
    assert abs(float(summary["D3"].value) - 225 / -655) < 1e-9
    assert summary["A4"].value == "办公费用"
    assert summary["A5"].value == "转账"
    assert summary["A6"].value == "合计"
    assert summary["B6"].value == 3
    assert money(summary["C6"]) == Decimal("-655.00")
    assert float(summary["D6"].value) == 1.0
    assert summary["A7"].value == "金额大写"
    assert summary["B7"].value == "负陆佰伍拾伍元整"


def test_company_layout_puts_description_before_company(connection, tmp_path) -> None:
    report, lines = build_report(connection)
    target = tmp_path / "july.xlsx"
    write_report_workbook(target, report, lines, layout_key="company")

    worksheet = load_workbook(target).active
    headers = [worksheet.cell(row=4, column=column).value for column in range(1, 9)]
    assert headers == [
        "序号",
        "日期",
        "费用名称",
        "金额",
        "报销人",
        "是否持有发票",
        "备注（报销费用说明）",
        "公司",
    ]
    assert worksheet["F5"].value == "有票"
    assert worksheet["G5"].value == "示例费用甲"
    assert worksheet["H5"].value == "示例公司"
    assert worksheet["G5"].alignment.wrap_text is True
    assert worksheet["H5"].alignment.wrap_text is None
    assert {str(rng) for rng in worksheet.merged_cells.ranges} == {
        "A1:G1",
        "A2:G2",
        "C3:G3",
        "A8:C8",
        "E8:H8",
    }


def test_plain_layout_has_seven_columns(connection, tmp_path) -> None:
    report, lines = build_report(connection)
    target = tmp_path / "plain.xlsx"
    write_report_workbook(target, report, lines, layout_key="plain")

    worksheet = load_workbook(target).active
    assert worksheet.max_column == 7
    assert worksheet.cell(row=4, column=7).value == "备注（报销费用说明）"
    assert worksheet.cell(row=5, column=7).value == "示例费用甲"
    assert {str(rng) for rng in worksheet.merged_cells.ranges} == {
        "A1:G1",
        "A2:G2",
        "C3:G3",
        "A8:C8",
        "E8:G8",
    }


def test_project_header_can_be_overridden(connection, tmp_path) -> None:
    report, lines = build_report(connection)
    target = tmp_path / "september.xlsx"
    write_report_workbook(
        target, report, lines, layout_key="project", project_header="费用发生项目"
    )
    worksheet = load_workbook(target).active
    assert worksheet.cell(row=4, column=7).value == "费用发生项目"


def test_invoice_note_is_preserved(connection, tmp_path) -> None:
    report_id = make_report(connection)
    add_line(connection, report_id, has_invoice=True, invoice_note="示例附注")
    report = reports.get_report(connection, report_id)
    lines = reports.report_lines(connection, report_id)
    target = tmp_path / "note.xlsx"
    write_report_workbook(target, report, lines)
    assert load_workbook(target).active["F5"].value == "有票（示例附注）"


def test_amount_mismatch_is_marked_in_the_sheet(connection, tmp_path) -> None:
    invoice_id = add_invoice(connection)
    report_id = make_report(connection)
    add_line(connection, report_id, amount_cents=10000, invoice_ids=[invoice_id])
    add_line(
        connection,
        report_id,
        description="示例费用乙",
        amount_cents=10000,
        has_invoice=True,
    )
    report = reports.get_report(connection, report_id)
    lines = reports.report_lines(connection, report_id)
    target = tmp_path / "mismatch.xlsx"
    write_report_workbook(target, report, lines)

    worksheet = load_workbook(target).active
    assert worksheet["D5"].font.bold is True
    assert worksheet["D5"].font.color.rgb.endswith("C00000")
    assert worksheet["D6"].font.bold is False


def test_empty_report_is_refused(connection, tmp_path) -> None:
    report_id = make_report(connection)
    report = reports.get_report(connection, report_id)
    with pytest.raises(ReportError, match="没有报销行"):
        write_report_workbook(tmp_path / "empty.xlsx", report, [])


def test_user_text_never_becomes_a_formula(connection, tmp_path) -> None:
    report_id = make_report(connection)
    add_line(
        connection,
        report_id,
        description='=HYPERLINK("http://example.invalid")',
        project_name="=1+1",
    )
    report = reports.get_report(connection, report_id)
    lines = reports.report_lines(connection, report_id)
    target = tmp_path / "formula.xlsx"
    write_report_workbook(target, report, lines)

    worksheet = load_workbook(target, data_only=False).active
    assert worksheet["G5"].value == "=1+1"
    assert worksheet["G5"].data_type == "s"
    assert worksheet["H5"].value == '=HYPERLINK("http://example.invalid")'
    assert worksheet["H5"].data_type == "s"


def test_long_description_is_truncated(connection, tmp_path) -> None:
    report_id = make_report(connection)
    add_line(connection, report_id, description="长" * 600)
    report = reports.get_report(connection, report_id)
    lines = reports.report_lines(connection, report_id)
    target = tmp_path / "long.xlsx"
    write_report_workbook(target, report, lines)

    value = load_workbook(target).active["H5"].value
    assert len(value) == 500
    assert value.endswith("…")


def test_writer_refuses_symlink_target(tmp_path, monkeypatch) -> None:
    connection = make_db(tmp_path)
    try:
        report, lines = build_report(connection)
    finally:
        connection.close()
    target = tmp_path / "target.xlsx"
    original = Path.is_symlink

    def fake_is_symlink(self):
        return self == target or original(self)

    monkeypatch.setattr(Path, "is_symlink", fake_is_symlink)
    with pytest.raises(ReportError, match="符号链接"):
        write_report_workbook(target, report, lines)


def test_report_filename_matches_company_convention(connection, tmp_path) -> None:
    report_id = reports.create_report(
        connection,
        company="示例公司",
        reporter="张三",
        month="2026-09",
        department="示例部门",
    )
    report = reports.get_report(connection, report_id)
    assert report_filename(report) == "示例公司费用报销支出明细表（张三）2609.xlsx"

    no_reporter = reports.create_report(
        connection, company="示例甲", reporter="李四", month="2026-07"
    )
    report = replace(reports.get_report(connection, no_reporter), reporter="")
    assert report_filename(report) == "示例甲费用报销支出明细表2607.xlsx"


def test_sanitize_filename_removes_invalid_characters() -> None:
    assert sanitize_filename("示例/测试:报销*单?") == "示例_测试_报销_单_"
    assert sanitize_filename("   ") == "报销单"
    assert sanitize_filename("名" * 200).endswith("名")
    assert len(sanitize_filename("名" * 200)) <= 120


def test_display_date_uses_paper_format() -> None:
    assert display_date("2026-09-01") == "2026.9.1"
    assert display_date("2026-10-10") == "2026.10.10"
    assert display_date("") == ""
    assert display_date("not-a-date") == "not-a-date"
