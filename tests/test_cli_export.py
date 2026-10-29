"""CLI tests for the export command: filters, guards, and error paths."""

from __future__ import annotations

import hashlib
from pathlib import Path

import fitz
import pytest
from openpyxl import load_workbook

from app.cli import _print_skipped, _unsafe_output_reason, main
from app.core.excel_export import EXCEL_FILENAME
from app.core.print_pdf import PRINT_FILENAME
from tests.pdf_factory import full_digital_lines, make_pdf

ROOT = Path(__file__).resolve().parents[1]


def invoice_lines(
    *,
    number: str = "26502000000000000001",
    buyer: str = "示例甲公司",
    date: str = "2026年07月30日",
    seller: str = "示例乙酒店",
    item: str = "*生产生活服务*住宿费",
    small: str = "¥ 1060.00",
    uppercase: str = "壹仟零陆拾圆整",
    net: str = "¥1000.00",
    tax: str = "¥60.00",
) -> list[tuple[float, float, str]]:
    lines = full_digital_lines(
        number=number,
        seller=seller,
        item=item,
        small=small,
        uppercase=uppercase,
        net=net,
        tax=tax,
    )
    replacements = {"示例甲公司": buyer, "2026年07月30日": date}
    return [(x, y, replacements.get(text, text)) for x, y, text in lines]


def write_invoice(folder: Path, name: str, **kwargs) -> Path:
    folder.mkdir(parents=True, exist_ok=True)
    return make_pdf(folder, invoice_lines(**kwargs), name=name)


def prepare_two_invoices(tmp_path: Path) -> tuple[Path, Path, list[Path]]:
    folder = tmp_path / "2026年7-8月"
    first = write_invoice(folder, "a.pdf")
    second = write_invoice(
        folder,
        "b.pdf",
        number="26502000000000000002",
        buyer="示例乙公司",
        date="2026年08月02日",
        seller="示例丙餐馆",
        item="*餐饮服务*餐费",
        small="¥ 88.00",
        uppercase="捌拾捌圆整",
        net="¥83.02",
        tax="¥4.98",
    )
    database = tmp_path / "pool.db"
    assert main(["import", str(folder), "--db", str(database)]) == 0
    return database, folder, [first, second]


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def test_export_end_to_end_writes_excel_and_pdf(tmp_path: Path, capsys) -> None:
    database, _, sources = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    hashes = {path: digest(path) for path in sources}
    output = tmp_path / "export"

    assert main(["export", "--db", str(database), "--out", str(output)]) == 0
    captured = capsys.readouterr()
    assert "已导出 2 张发票" in captured.out
    assert "合计 1,148.00 元" in captured.out
    assert "1 页" in captured.out
    assert EXCEL_FILENAME in captured.out
    assert (output / EXCEL_FILENAME).exists()

    workbook = load_workbook(output / EXCEL_FILENAME)
    assert workbook.sheetnames == ["发票明细", "分类汇总", "公司汇总", "导出信息"]
    detail = workbook["发票明细"]
    names = {detail["O2"].value, detail["O3"].value}
    assert names == {"a.pdf", "b.pdf"}
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert str(tmp_path) not in cell.value

    categories = workbook["分类汇总"]
    assert categories["A2"].value == "差旅费"
    assert categories["B2"].value == 1
    assert float(categories["C2"].value) == 1000.0
    assert float(categories["D2"].value) == 60.0
    assert float(categories["E2"].value) == 1060.0
    assert categories["A3"].value == "餐饮费"
    assert float(categories["E3"].value) == 88.0
    assert categories["B4"].value == 2
    assert float(categories["E4"].value) == 1148.0
    assert float(categories["F4"].value) == 1.0

    buyers = workbook["公司汇总"]
    assert buyers["A2"].value == "示例甲公司"
    assert buyers["B2"].value == 1
    assert buyers["C2"].value == 0
    assert float(buyers["D2"].value) == 1060.0
    assert buyers["A3"].value == "示例乙公司"
    assert float(buyers["D3"].value) == 88.0
    assert buyers["B4"].value == 2
    assert float(buyers["D4"].value) == 1148.0

    meta = {
        row[0].value: row[1].value
        for row in workbook["导出信息"].iter_rows(min_row=2)
        if row[0].value and row[1].value is not None
    }
    assert meta["筛选条件"] == "全部发票"
    assert meta["发票条数"] == 2
    assert float(meta["价税合计"]) == 1148.0
    assert meta["导出时间"]

    document = fitz.open(output / PRINT_FILENAME)
    try:
        assert document.page_count == 1
        text = document[0].get_text()
        assert "a.pdf p1/1" in text
        assert "b.pdf p1/1" in text
        assert str(tmp_path) not in text
    finally:
        document.close()

    group_dir = output / "发票分组"
    assert (group_dir / "示例甲公司" / "差旅费.pdf").is_file()
    assert (group_dir / "示例乙公司" / "餐饮费.pdf").is_file()
    assert "分类 PDF" in captured.out
    assert "示例甲公司 / 差旅费" in captured.out
    grouped = fitz.open(group_dir / "示例甲公司" / "差旅费.pdf")
    try:
        # Clean merge: only the invoice text, no sequence/label header.
        assert "示例乙酒店" in grouped[0].get_text()
        assert "a.pdf p1" not in grouped[0].get_text()
    finally:
        grouped.close()

    for path, before in hashes.items():
        assert digest(path) == before


def test_export_filters_by_buyer_date_and_batch(tmp_path: Path, capsys) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    september = tmp_path / "2026年9月"
    write_invoice(
        september,
        "c.pdf",
        number="26502000000000000003",
        buyer="示例丙公司",
        date="2026年09月03日",
        seller="示例丁酒店",
    )
    assert main(["import", str(september), "--db", str(database)]) == 0
    capsys.readouterr()

    buyer_out = tmp_path / "by_buyer"
    assert (
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(buyer_out),
                "--buyer",
                "示例乙公司",
            ]
        )
        == 0
    )
    captured = capsys.readouterr()
    assert "已导出 1 张发票" in captured.out
    workbook = load_workbook(buyer_out / EXCEL_FILENAME)
    assert workbook["发票明细"]["O2"].value == "b.pdf"

    date_out = tmp_path / "by_date"
    assert (
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(date_out),
                "--from",
                "2026-08-01",
                "--to",
                "2026-08-31",
            ]
        )
        == 0
    )
    workbook = load_workbook(date_out / EXCEL_FILENAME)
    assert workbook["发票明细"]["O2"].value == "b.pdf"

    batch_out = tmp_path / "by_batch"
    assert (
        main(
            ["export", "--db", str(database), "--out", str(batch_out), "--batch", "1"]
        )
        == 0
    )
    workbook = load_workbook(batch_out / EXCEL_FILENAME)
    assert {workbook["发票明细"]["O2"].value, workbook["发票明细"]["O3"].value} == {
        "a.pdf",
        "b.pdf",
    }
    assert workbook["发票明细"]["O4"].value is None

    second_batch = tmp_path / "second_batch"
    assert (
        main(
            ["export", "--db", str(database), "--out", str(second_batch), "--batch", "2"]
        )
        == 0
    )
    workbook = load_workbook(second_batch / EXCEL_FILENAME)
    assert workbook["发票明细"]["O2"].value == "c.pdf"
    assert workbook["发票明细"]["O3"].value is None

    capsys.readouterr()
    assert (
        main(
            ["export", "--db", str(database), "--out", str(tmp_path / "nope"), "--batch", "9"]
        )
        == 1
    )
    assert "批次不存在" in capsys.readouterr().err


def test_export_excludes_review_invoices(tmp_path: Path, capsys) -> None:
    folder = tmp_path / "july"
    write_invoice(folder, "good.pdf")
    (folder / "broken.pdf").write_bytes(b"not a real pdf")
    database = tmp_path / "pool.db"
    assert main(["import", str(folder), "--db", str(database)]) == 0
    capsys.readouterr()

    all_out = tmp_path / "all"
    assert main(["export", "--db", str(database), "--out", str(all_out)]) == 0
    workbook = load_workbook(all_out / EXCEL_FILENAME)
    assert workbook["发票明细"]["O2"].value == "good.pdf"
    assert workbook["发票明细"]["L2"].value is None
    assert workbook["发票明细"]["O3"].value == "broken.pdf"
    assert workbook["发票明细"]["L3"].value == "是"
    assert "待确认 1 张" in capsys.readouterr().out

    clean_out = tmp_path / "clean"
    assert (
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(clean_out),
                "--exclude-review",
            ]
        )
        == 0
    )
    captured = capsys.readouterr()
    assert "已导出 1 张发票" in captured.out
    workbook = load_workbook(clean_out / EXCEL_FILENAME)
    assert workbook["发票明细"]["O2"].value == "good.pdf"
    assert workbook["发票明细"]["O3"].value is None


def test_export_without_matching_rows_returns_1(tmp_path: Path, capsys) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    output = tmp_path / "empty"
    assert (
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(output),
                "--buyer",
                "不存在公司",
            ]
        )
        == 1
    )
    assert "没有符合条件的发票" in capsys.readouterr().out
    assert not output.exists()


def test_export_missing_database_returns_1(tmp_path: Path, capsys) -> None:
    assert (
        main(["export", "--db", str(tmp_path / "missing.db"), "--out", str(tmp_path / "out")])
        == 1
    )
    assert "数据库不存在" in capsys.readouterr().err


def test_export_rejects_reversed_or_bad_dates(tmp_path: Path, capsys) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    assert (
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(tmp_path / "out"),
                "--from",
                "2026-09-01",
                "--to",
                "2026-08-01",
            ]
        )
        == 2
    )
    assert "--from 不能晚于 --to" in capsys.readouterr().err

    with pytest.raises(SystemExit):
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(tmp_path / "out"),
                "--from",
                "2026/08/01",
            ]
        )


def test_export_refuses_repository_output(
    tmp_path: Path, capsys, monkeypatch
) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    fake_root = tmp_path / "repo"
    fake_root.mkdir()
    monkeypatch.setattr("app.core.paths.repo_root", lambda: fake_root)
    target = fake_root / "export"
    assert main(["export", "--db", str(database), "--out", str(target)]) == 2
    assert "输出路径在仓库内非忽略目录" in capsys.readouterr().err
    assert not target.exists()


def test_export_rejects_dotdot_repository_output(
    tmp_path: Path, capsys, monkeypatch
) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    fake_root = tmp_path / "repo"
    fake_root.mkdir()
    monkeypatch.setattr("app.core.paths.repo_root", lambda: fake_root)
    target = fake_root / "out" / ".." / "escaped"
    assert main(["export", "--db", str(database), "--out", str(target)]) == 2
    assert "输出路径在仓库内非忽略目录" in capsys.readouterr().err
    assert not (fake_root / "escaped").exists()


def test_export_force_allows_repository_output(
    tmp_path: Path, capsys, monkeypatch
) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    fake_root = tmp_path / "repo"
    fake_root.mkdir()
    monkeypatch.setattr("app.core.paths.repo_root", lambda: fake_root)
    target = fake_root / "export"
    assert (
        main(
            [
                "export",
                "--db",
                str(database),
                "--out",
                str(target),
                "--force",
            ]
        )
        == 0
    )
    assert (target / EXCEL_FILENAME).exists()
    assert (target / PRINT_FILENAME).exists()


def test_export_rejects_file_as_output(tmp_path: Path, capsys) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    target = tmp_path / "not_a_directory.txt"
    target.write_text("keep me", encoding="utf-8")
    assert main(["export", "--db", str(database), "--out", str(target)]) == 2
    assert "输出路径不是目录" in capsys.readouterr().err
    assert target.read_text(encoding="utf-8") == "keep me"


def test_output_guard_reasons() -> None:
    assert _unsafe_output_reason(ROOT / "out" / "export") is None
    assert _unsafe_output_reason(ROOT / "data" / "export") is None
    assert _unsafe_output_reason(ROOT / "_test_export_guard") == "在仓库内非忽略目录"
    assert _unsafe_output_reason(Path(r"\\server\share\export")) == "位于网络共享路径"


def test_export_refuses_remote_output(
    tmp_path: Path, capsys, monkeypatch
) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    monkeypatch.setattr("app.cli.remote_storage_reason", lambda path: "网络盘")
    target = tmp_path / "remote_out"
    assert main(["export", "--db", str(database), "--out", str(target)]) == 2
    assert "输出路径位于网络盘" in capsys.readouterr().err
    assert not target.exists()


def test_export_force_allows_remote_output(
    tmp_path: Path, capsys, monkeypatch
) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    capsys.readouterr()
    monkeypatch.setattr("app.cli.remote_storage_reason", lambda path: "OneDrive")
    target = tmp_path / "remote_force"
    assert (
        main(["export", "--db", str(database), "--out", str(target), "--force"])
        == 0
    )
    assert (target / EXCEL_FILENAME).exists()
    assert (target / PRINT_FILENAME).exists()


def test_output_guard_reports_remote_before_repo(
    monkeypatch,
) -> None:
    monkeypatch.setattr("app.cli.remote_storage_reason", lambda path: "OneDrive")
    assert _unsafe_output_reason(ROOT / "_test_export_guard") == "位于OneDrive"


def test_export_missing_source_still_writes_excel(tmp_path: Path, capsys) -> None:
    folder = tmp_path / "july"
    source = write_invoice(folder, "gone.pdf")
    database = tmp_path / "pool.db"
    assert main(["import", str(folder), "--db", str(database)]) == 0
    source.unlink()
    capsys.readouterr()

    output = tmp_path / "export"
    assert main(["export", "--db", str(database), "--out", str(output)]) == 0
    captured = capsys.readouterr()
    assert "已导出 1 张发票" in captured.out
    assert "缺少源文件 1 张" in captured.err
    assert "已生成说明页" in captured.err
    assert (output / EXCEL_FILENAME).exists()
    document = fitz.open(output / PRINT_FILENAME)
    try:
        assert "本次没有可打印的票面" in document[0].get_text()
    finally:
        document.close()


def test_export_partial_missing_source_reports_gap(tmp_path: Path, capsys) -> None:
    folder = tmp_path / "july"
    write_invoice(folder, "a.pdf")
    missing = write_invoice(
        folder,
        "b.pdf",
        number="26502000000000000002",
        buyer="示例乙公司",
    )
    database = tmp_path / "pool.db"
    assert main(["import", str(folder), "--db", str(database)]) == 0
    missing.unlink()
    capsys.readouterr()

    output = tmp_path / "export"
    assert main(["export", "--db", str(database), "--out", str(output)]) == 0
    captured = capsys.readouterr()
    assert "1 页，1 张票面，1 条记录" in captured.out
    assert "缺少源文件 1 张" in captured.err
    assert "打印包缺少 1 条发票记录" in captured.err

    document = fitz.open(output / PRINT_FILENAME)
    try:
        text = document[0].get_text()
        assert "a.pdf p1/1" in text
        assert "b.pdf" not in text
    finally:
        document.close()


def test_print_skipped_truncates_long_lists(capsys) -> None:
    _print_skipped("缺少源文件", [f"f{index}.pdf" for index in range(7)])
    error = capsys.readouterr().err
    assert "警告: 缺少源文件 7 张：" in error
    assert "f0.pdf" in error
    assert "f4.pdf" in error
    assert "f5.pdf" not in error
    assert error.rstrip().endswith("…")

    _print_skipped("缺少源文件", [])
    assert capsys.readouterr().err == ""


def test_export_write_failure_is_reported_without_paths(
    tmp_path: Path, capsys, monkeypatch
) -> None:
    database, _, _ = prepare_two_invoices(tmp_path)
    output = tmp_path / "out"
    assert main(["export", "--db", str(database), "--out", str(output)]) == 0
    capsys.readouterr()
    before = {
        name: digest(output / name) for name in (EXCEL_FILENAME, PRINT_FILENAME)
    }

    def boom(*args, **kwargs):
        raise OSError(r"C:\secret\export.xlsx")

    monkeypatch.setattr("app.cli.write_workbook", boom)
    assert main(["export", "--db", str(database), "--out", str(output)]) == 2
    error = capsys.readouterr().err
    assert "无法写入导出文件" in error
    assert "secret" not in error
    for name, original in before.items():
        assert digest(output / name) == original
    assert list(output.glob("*.tmp")) == []
