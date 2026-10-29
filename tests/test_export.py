"""Tests for the Excel package and the 2-up A4 print PDF."""

from __future__ import annotations

import hashlib
import itertools
import shutil
from decimal import Decimal
from pathlib import Path

import fitz
import pytest
from openpyxl import Workbook, load_workbook

from app.core import db, excel_export
from app.core.excel_export import DETAIL_HEADERS, write_workbook
from app.core.export import (
    ExportError,
    ExportFilters,
    ExportRow,
    collect_rows,
    summarize,
)
from app.core.models import InvoiceType, ParsedInvoice
from app.core.print_pdf import write_print_pdf
from tests.pdf_factory import full_digital_lines, make_pdf, make_pdf_pages

_counter = itertools.count(1)


@pytest.mark.parametrize(
    "value",
    ["=1+1", "+1+1", "-1+1", "@SUM(A1)", "\ttext", "'=text", "=cmd|calc"],
)
def test_text_cells_are_always_written_as_text(value: str) -> None:
    workbook = Workbook()
    cell = workbook.active["A1"]
    excel_export._set_text(cell, value)
    assert cell.data_type == "s"


def make_row(**overrides) -> ExportRow:
    values = {
        "id": 1,
        "batch_id": 1,
        "source_path": r"D:\invoices\a.pdf",
        "page_index": 0,
        "page_count": 1,
        "invoice_type": InvoiceType.FULL_DIGITAL_NORMAL.value,
        "invoice_number": "26502000000000000001",
        "invoice_date": "2026-07-30",
        "seller_name": "示例乙酒店",
        "buyer_name": "示例甲公司",
        "amount_net_cents": 100000,
        "tax_amount_cents": 6000,
        "amount_total_cents": 106000,
        "item_summary": "住宿费",
        "category": "差旅费",
        "confidence": {},
        "warnings": (),
        "needs_review": False,
    }
    values.update(overrides)
    return ExportRow(**values)


def add_invoice(
    connection,
    *,
    source_path: Path,
    batch_id: int | None = None,
    number: str | None = None,
    date: str = "2026-07-30",
    seller: str = "示例乙酒店",
    buyer: str = "示例甲公司",
    net: str | None = "1000.00",
    tax: str | None = "60.00",
    total: str | None = "1060.00",
    item: str = "住宿费",
    category: str = "差旅费",
    warnings: tuple[str, ...] = (),
    page_index: int = 0,
    page_count: int = 1,
) -> int:
    index = next(_counter)
    if number is None:
        number = f"26502000000000{index:06d}"
    if batch_id is None:
        batch_id = db.create_batch(
            connection,
            label=f"批次{index}",
            source_dir=str(source_path.parent),
        )
    sha = hashlib.sha256(f"{source_path}|{page_index}|{index}".encode()).hexdigest()
    db.insert_file(
        connection,
        sha256=sha,
        batch_id=batch_id,
        source_path=str(source_path),
        page_count=page_count,
    )
    invoice = ParsedInvoice(
        source_path=str(source_path),
        page_index=page_index,
        page_count=page_count,
        invoice_type=InvoiceType.FULL_DIGITAL_NORMAL,
        invoice_number=number,
        invoice_date=date,
        seller_name=seller,
        seller_tax_id="92500000TEST00002X",
        buyer_name=buyer,
        buyer_tax_id="91500000TEST00001X",
        amount_total=None if total is None else Decimal(total),
        amount_net=None if net is None else Decimal(net),
        tax_amount=None if tax is None else Decimal(tax),
        item_summary=item,
        category=category,
        category_rule="test",
        confidence={},
        warnings=list(warnings),
    )
    db.insert_invoice(connection, batch_id=batch_id, sha256=sha, invoice=invoice)
    connection.commit()
    return batch_id


def test_collect_rows_filters_and_summary(tmp_path: Path) -> None:
    connection = db.connect(tmp_path / "pool.db")
    try:
        batch1 = add_invoice(
            connection, source_path=tmp_path / "a.pdf", total="100.00"
        )
        add_invoice(
            connection,
            source_path=tmp_path / "b.pdf",
            batch_id=batch1,
            date="2026-07-15",
            buyer="示例乙公司",
            total="200.00",
            warnings=("测试警告",),
        )
        batch2 = add_invoice(
            connection,
            source_path=tmp_path / "c.pdf",
            date="2026-09-01",
            buyer="示例甲公司",
            total="300.00",
        )
        add_invoice(
            connection,
            source_path=tmp_path / "d.pdf",
            batch_id=batch2,
            date="",
            buyer="",
            total="400.00",
        )

        rows = collect_rows(connection, ExportFilters())
        assert [row.invoice_date for row in rows] == [
            "2026-07-15",
            "2026-07-30",
            "2026-09-01",
            "",
        ]
        summary = summarize(rows)
        assert summary.count == 4
        assert summary.review_count == 2  # warning row + undated row
        assert summary.missing_amount_count == 0
        assert summary.total_cents == 100000  # 100 + 200 + 300 + 400 yuan

        assert len(collect_rows(connection, ExportFilters(batch_id=batch1))) == 2
        assert len(collect_rows(connection, ExportFilters(buyer="示例甲公司"))) == 2
        assert (
            len(collect_rows(connection, ExportFilters(buyer=" 示例甲公司 "))) == 2
        )
        ranged = collect_rows(
            connection,
            ExportFilters(date_from="2026-08-01", date_to="2026-09-30"),
        )
        assert [row.source_name for row in ranged] == ["c.pdf"]
        assert [
            row.source_name
            for row in collect_rows(
                connection,
                ExportFilters(date_from="2026-07-15", date_to="2026-07-15"),
            )
        ] == ["b.pdf"]  # boundaries are inclusive
        assert [
            row.source_name
            for row in collect_rows(connection, ExportFilters(date_from="2026-09-01"))
        ] == ["c.pdf"]
        assert [
            row.source_name
            for row in collect_rows(connection, ExportFilters(date_to="2026-07-15"))
        ] == ["b.pdf"]
        undated_excluded = collect_rows(
            connection, ExportFilters(date_from="1990-01-01")
        )
        assert "d.pdf" not in {row.source_name for row in undated_excluded}
        assert "d.pdf" in {row.source_name for row in collect_rows(connection, ExportFilters())}
        kept = collect_rows(connection, ExportFilters(include_review=False))
        assert [row.source_name for row in kept] == ["a.pdf", "c.pdf"]
        assert collect_rows(
            connection, ExportFilters(buyer="示例乙公司", include_review=False)
        ) == []
    finally:
        connection.close()


def test_filter_description_is_explicit() -> None:
    assert ExportFilters().describe() == "全部发票"
    filters = ExportFilters(
        batch_id=2,
        buyer="示例甲公司",
        date_from="2026-07-01",
        date_to="2026-07-31",
        include_review=False,
    )
    assert filters.describe() == (
        "批次 #2；购买方「示例甲公司」；"
        "开票日期 2026-07-01 至 2026-07-31；排除待确认"
    )
    assert filters.review_mode() == "已排除"
    assert ExportFilters().review_mode() == "包含，并在明细中标记"
    assert (
        ExportFilters(date_from="2026-08-01").describe()
        == "开票日期 ≥ 2026-08-01"
    )
    assert (
        ExportFilters(date_to="2026-08-31").describe()
        == "开票日期 ≤ 2026-08-31"
    )


def test_export_row_traceability_fields() -> None:
    row = make_row(
        source_path=r"D:\secret\invoices\发票A.pdf",
        page_index=1,
        page_count=3,
        invoice_type=InvoiceType.TRAIN.value,
        confidence={
            "seller_name": "medium",
            "category": "low",
            "unknown_field": "low",
            "buyer_name": "high",
        },
        warnings=(
            "甲",
            r"无法读取 D:\secret\invoice.pdf",
            r"\\server\share\x.pdf",
            r"无法读取 C:\Users\张 三\inv.pdf",
            r"\\nas\财务 共享\a.pdf",
            "C:/Users/zhang/a.pdf",
            "/home/user/我的 文档/a.pdf",
            "乙",
        ),
        needs_review=True,
    )
    assert row.source_name == "发票A.pdf"
    assert row.page_label == "2/3"
    assert row.type_label == "铁路客票"
    assert row.confidence_notes == "销售方(中)；分类(低)"
    assert row.warning_text == (
        "甲；无法读取 <路径已隐藏>；<路径已隐藏>；无法读取 <路径已隐藏>；"
        "<路径已隐藏>；<路径已隐藏>；<路径已隐藏>；乙"
    )
    for leaked in ("secret", "server", "张 三", "财务 共享", "zhang", "我的 文档"):
        assert leaked not in row.warning_text


def test_workbook_sheets_values_and_no_leaked_paths(tmp_path: Path) -> None:
    rows = [
        make_row(
            id=1,
            source_path=str(tmp_path / "secret-dir" / "发票A.pdf"),
            invoice_number="26502000000000000001",
            confidence={"seller_name": "medium", "category": "high"},
        ),
        make_row(
            id=2,
            source_path=str(tmp_path / "other" / "b.pdf"),
            page_index=1,
            page_count=2,
            invoice_number="26502000000000000002",
            seller_name='=HYPERLINK("http://example.invalid")',
            item_summary="打印纸",
            category="办公费用",
            amount_net_cents=None,
            tax_amount_cents=None,
            amount_total_cents=5000,
            confidence={"category": "low", "amount_total": "low"},
            warnings=("测试警告",),
            needs_review=True,
        ),
        make_row(
            id=3,
            source_path=str(tmp_path / "c.pdf"),
            invoice_number="26502000000000000003",
            buyer_name="示例乙公司",
            amount_net_cents=8000,
            tax_amount_cents=1000,
            amount_total_cents=9000,
        ),
    ]
    summary = summarize(rows)
    target = tmp_path / "out" / "发票明细.xlsx"
    write_workbook(
        target,
        rows,
        summary,
        filters=ExportFilters(),
        generated_at="2026-10-28 12:00:00",
    )

    workbook = load_workbook(target, data_only=False)
    assert workbook.sheetnames == ["发票明细", "分类汇总", "公司汇总", "导出信息"]

    detail = workbook["发票明细"]
    assert [cell.value for cell in detail[1]] == list(DETAIL_HEADERS)
    assert detail["A2"].value == 1
    assert detail["B2"].value == "2026-07-30"
    assert detail["C2"].value == "26502000000000000001"
    assert detail["D2"].value == "全电普票"
    assert detail["H2"].value == "差旅费"
    assert float(detail["I2"].value) == 1000.0
    assert float(detail["J2"].value) == 60.0
    assert float(detail["K2"].value) == 1060.0
    assert detail["M2"].value == "销售方(中)"
    assert detail["O2"].value == "发票A.pdf"
    assert detail["P2"].value == "1/1"

    assert detail["L3"].value == "是"
    assert detail["M3"].value == "金额(低)；分类(低)"
    assert "测试警告" in detail["N3"].value
    assert detail["O3"].value == "b.pdf"
    assert detail["P3"].value == "2/2"
    assert detail["I3"].value is None
    assert detail["E3"].value == '=HYPERLINK("http://example.invalid")'
    assert detail["E3"].data_type == "s"  # must stay text, not a formula

    assert str(detail["H6"].value).startswith("合计")
    assert float(detail["I6"].value) == 1080.0
    assert float(detail["J6"].value) == 70.0
    assert float(detail["K6"].value) == 1200.0
    assert detail["G7"].value == "价税合计（大写）"
    assert detail["H7"].value == "壹仟贰佰元整"
    assert detail.auto_filter.ref == "A1:P4"

    categories = workbook["分类汇总"]
    assert categories["A2"].value == "差旅费"
    assert categories["B2"].value == 2
    assert float(categories["E2"].value) == 1150.0
    assert abs(float(categories["F2"].value) - 1150 / 1200) < 1e-9
    assert categories["A3"].value == "办公费用"
    assert float(categories["E3"].value) == 50.0
    assert float(categories["E4"].value) == 1200.0
    assert categories["A5"].value == "价税合计（大写）"
    assert categories["B5"].value == "壹仟贰佰元整"

    buyers = workbook["公司汇总"]
    assert buyers["A2"].value == "示例甲公司"
    assert buyers["B2"].value == 2
    assert buyers["C2"].value == 1
    assert float(buyers["D2"].value) == 1110.0
    assert buyers["A3"].value == "示例乙公司"
    assert float(buyers["D3"].value) == 90.0

    meta = {
        row[0].value: row[1].value
        for row in workbook["导出信息"].iter_rows(min_row=2)
        if row[0].value and row[1].value is not None
    }
    assert meta["筛选条件"] == "全部发票"
    assert meta["待确认处理"] == "包含，并在明细中标记"
    assert meta["发票条数"] == 3
    assert meta["待确认条数"] == 1
    assert float(meta["价税合计"]) == 1200.0
    assert meta["价税合计（大写）"] == "壹仟贰佰元整"
    assert meta["来源票面数"] == 3

    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows():
            for cell in row:
                if isinstance(cell.value, str):
                    assert "secret-dir" not in cell.value
                    assert str(tmp_path) not in cell.value


def test_print_pdf_two_up_layout_and_sources_untouched(tmp_path: Path) -> None:
    pdf1 = make_pdf(tmp_path, full_digital_lines(), name="a.pdf")
    pdf2 = make_pdf_pages(
        tmp_path,
        [
            full_digital_lines(number="26502000000000000002"),
            full_digital_lines(number="26502000000000000003"),
        ],
        name="b.pdf",
    )
    hashes = {
        path: hashlib.sha256(path.read_bytes()).hexdigest() for path in (pdf1, pdf2)
    }
    rows = [
        make_row(id=1, source_path=str(pdf1), invoice_number="26502000000000000001"),
        make_row(
            id=2,
            source_path=str(pdf2),
            invoice_number="26502000000000000002",
            page_count=2,
        ),
        make_row(
            id=3,
            source_path=str(pdf2),
            page_index=1,
            page_count=2,
            invoice_number="26502000000000000003",
        ),
    ]

    target = tmp_path / "out" / "print.pdf"
    info = write_print_pdf(target, rows)
    assert info.pages == 2
    assert info.unique_pages == 3
    assert info.records == 3
    assert info.duplicate_records == 0
    assert info.skipped_missing == ()
    assert info.skipped_invalid == ()

    document = fitz.open(target)
    try:
        assert document.page_count == 2
        for page in document:
            assert round(page.rect.width, 2) == 595.28
            assert round(page.rect.height, 2) == 841.89
        first = document[0].get_text()
        second = document[1].get_text()
        assert "a.pdf p1/1" in first
        assert "b.pdf p1/2" in first
        assert "b.pdf p2/2" not in first
        assert "b.pdf p2/2" in second
        assert "3. 2026-07-30" in second
    finally:
        document.close()

    for path, digest in hashes.items():
        assert hashlib.sha256(path.read_bytes()).hexdigest() == digest


def test_print_pdf_dedups_shared_source_page(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path, full_digital_lines(), name="shared.pdf")
    rows = [
        make_row(id=1, source_path=str(pdf), needs_review=False),
        make_row(id=2, source_path=str(pdf), needs_review=True),
    ]
    target = tmp_path / "print.pdf"
    info = write_print_pdf(target, rows)
    assert info.pages == 1
    assert info.unique_pages == 1
    assert info.records == 2
    assert info.duplicate_records == 1

    document = fitz.open(target)
    try:
        text = document[0].get_text()
        assert text.count("shared.pdf p1/1") == 1
        assert "本页 2 张" in text
        assert "待确认" in text
    finally:
        document.close()


def test_print_pdf_reports_missing_and_corrupt_sources(tmp_path: Path) -> None:
    good = make_pdf(tmp_path, full_digital_lines(), name="good.pdf")
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a real pdf")
    rows = [
        make_row(id=1, source_path=str(good)),
        make_row(id=2, source_path=str(tmp_path / "missing.pdf")),
        make_row(id=3, source_path=str(broken)),
    ]
    target = tmp_path / "print.pdf"
    info = write_print_pdf(target, rows)
    assert info.pages == 1
    assert info.unique_pages == 1
    assert info.records == 1
    assert info.skipped_missing == ("missing.pdf p1",)
    assert info.skipped_invalid == ("broken.pdf p1",)
    assert target.exists()


def test_print_pdf_without_cjk_font_keeps_ascii_label(
    tmp_path: Path, monkeypatch
) -> None:
    from app.core import print_pdf

    pdf = make_pdf(tmp_path, full_digital_lines(), name="invoice.pdf")
    monkeypatch.setattr(
        print_pdf,
        "_fonts",
        lambda: print_pdf._Fonts(latin=fitz.Font("helv"), cjk=None),
    )
    target = tmp_path / "print.pdf"
    info = write_print_pdf(target, [make_row(source_path=str(pdf))])
    assert info.pages == 1
    assert info.skipped_invalid == ()
    document = fitz.open(target)
    try:
        text = document[0].get_text()
        assert "2026-07-30" in text
        assert "invoice.pdf p1/1" in text
    finally:
        document.close()


def test_print_pdf_second_slot_failure_keeps_first(
    tmp_path: Path, monkeypatch
) -> None:
    first = make_pdf(tmp_path, full_digital_lines(), name="first.pdf")
    second = make_pdf(
        tmp_path,
        full_digital_lines(number="26502000000000000002"),
        name="second.pdf",
    )
    original = fitz.Page.show_pdf_page
    source_calls = {"count": 0}

    def flaky(page, rect, docsrc, pno, *args, **kwargs):
        # Label overlays are 20pt tall; invoice source pages are not.
        if docsrc[0].rect.height > 100:
            source_calls["count"] += 1
            if source_calls["count"] == 2:
                raise RuntimeError("simulated render failure")
        return original(page, rect, docsrc, pno, *args, **kwargs)

    monkeypatch.setattr(fitz.Page, "show_pdf_page", flaky)
    target = tmp_path / "print.pdf"
    info = write_print_pdf(
        target,
        [
            make_row(id=1, source_path=str(first)),
            make_row(id=2, source_path=str(second), invoice_number="26502000000000000002"),
        ],
    )
    assert info.pages == 1
    assert info.unique_pages == 1
    assert info.records == 1
    assert info.skipped_invalid == ("second.pdf p1",)

    document = fitz.open(target)
    try:
        assert document.page_count == 1
        text = document[0].get_text()
        assert "first.pdf p1/1" in text
        assert "second.pdf" not in text
    finally:
        document.close()


def test_writers_refuse_symlink_targets(tmp_path: Path, monkeypatch) -> None:
    rows = [make_row(source_path=str(tmp_path / "a.pdf"))]
    summary = summarize(rows)
    excel_target = tmp_path / "target.xlsx"
    pdf_target = tmp_path / "target.pdf"
    original = Path.is_symlink

    def fake_is_symlink(self):
        return self in {excel_target, pdf_target} or original(self)

    monkeypatch.setattr(Path, "is_symlink", fake_is_symlink)
    with pytest.raises(ExportError):
        write_workbook(
            excel_target,
            rows,
            summary,
            filters=ExportFilters(),
            generated_at="2026-10-28 12:00:00",
        )
    with pytest.raises(ExportError):
        write_print_pdf(pdf_target, rows)


def test_print_pdf_skips_encrypted_and_out_of_range_pages(tmp_path: Path) -> None:
    plain = make_pdf(tmp_path, full_digital_lines(), name="plain.pdf")
    encrypted = tmp_path / "encrypted.pdf"
    protection = fitz.open()
    protection.new_page()
    protection.save(
        str(encrypted),
        encryption=fitz.PDF_ENCRYPT_AES_256,
        user_pw="secret",
    )
    protection.close()

    rows = [
        make_row(id=1, source_path=str(plain)),
        make_row(id=2, source_path=str(encrypted)),
        make_row(id=3, source_path=str(plain), page_index=5, page_count=1),
    ]
    target = tmp_path / "print.pdf"
    info = write_print_pdf(target, rows)
    assert info.pages == 1
    assert info.unique_pages == 1
    assert info.records == 1
    assert "encrypted.pdf p1" in info.skipped_invalid
    assert "plain.pdf p6" in info.skipped_invalid


def test_print_pdf_with_no_printable_page_writes_notice(tmp_path: Path) -> None:
    rows = [make_row(source_path=str(tmp_path / "missing.pdf"))]
    target = tmp_path / "print.pdf"
    info = write_print_pdf(target, rows)
    assert info.pages == 0
    assert info.unique_pages == 0
    assert info.notice_only is True
    assert info.skipped_missing == ("missing.pdf p1",)
    assert target.exists()

    document = fitz.open(target)
    try:
        text = document[0].get_text()
        assert "本次没有可打印的票面" in text
        assert "缺少源文件 1 条" in text
    finally:
        document.close()


def test_print_pdf_overwrites_stale_previous_export(tmp_path: Path) -> None:
    pdf = make_pdf(tmp_path, full_digital_lines(), name="old.pdf")
    target = tmp_path / "print.pdf"
    write_print_pdf(target, [make_row(source_path=str(pdf))])
    first = fitz.open(target)
    assert "old.pdf p1/1" in first[0].get_text()
    first.close()

    info = write_print_pdf(
        target, [make_row(source_path=str(tmp_path / "missing.pdf"))]
    )
    assert info.notice_only is True
    document = fitz.open(target)
    try:
        text = document[0].get_text()
        assert "本次没有可打印的票面" in text
        assert "old.pdf" not in text
    finally:
        document.close()


def test_print_pdf_embeds_more_sources_than_windows_stdio_limit(
    tmp_path: Path,
) -> None:
    """Regression: one OS handle per embedded source exhausted the 512-file
    CRT limit on Windows and the save failed with ``Too many open files``."""
    base = make_pdf(tmp_path, full_digital_lines(), name="base.pdf")
    rows = []
    for index in range(520):
        copy = tmp_path / f"copy-{index:03d}.pdf"
        shutil.copyfile(base, copy)
        rows.append(
            make_row(
                id=index + 1,
                source_path=str(copy),
                invoice_number=f"26502000000000{index:06d}",
            )
        )

    info = write_print_pdf(tmp_path / "print.pdf", rows)
    assert info.records == 520
    assert info.unique_pages == 520
    assert info.skipped_missing == ()
    assert info.skipped_invalid == ()
    assert info.pages == 260


def test_collect_rows_tolerates_bad_json(tmp_path: Path) -> None:
    connection = db.connect(tmp_path / "pool.db")
    try:
        add_invoice(connection, source_path=tmp_path / "a.pdf")
        connection.execute(
            "UPDATE invoices SET warnings = 'not json', confidence = '[1,2]'"
        )
        connection.commit()
        rows = collect_rows(connection, ExportFilters())
        assert len(rows) == 1
        assert rows[0].warning_text == ""
        assert rows[0].confidence_notes == ""
    finally:
        connection.close()


def test_summary_falls_back_for_unknown_buyer_and_category(tmp_path: Path) -> None:
    rows = [
        make_row(
            id=1,
            source_path=str(tmp_path / "a.pdf"),
            buyer_name="",
            category="",
            amount_net_cents=None,
            tax_amount_cents=None,
            amount_total_cents=1000,
        )
    ]
    target = tmp_path / "fallback.xlsx"
    write_workbook(
        target,
        rows,
        summarize(rows),
        filters=ExportFilters(),
        generated_at="2026-10-28 12:00:00",
    )
    workbook = load_workbook(target)
    assert workbook["分类汇总"]["A2"].value == "未分类"
    assert workbook["公司汇总"]["A2"].value == "未识别"


def test_workbook_handles_negative_amounts_and_zero_total(tmp_path: Path) -> None:
    rows = [
        make_row(
            id=1,
            source_path=str(tmp_path / "plus.pdf"),
            category="差旅费",
            amount_net_cents=100000,
            tax_amount_cents=6000,
            amount_total_cents=106000,
        ),
        make_row(
            id=2,
            source_path=str(tmp_path / "minus.pdf"),
            category="其他费用",
            amount_net_cents=-100000,
            tax_amount_cents=-6000,
            amount_total_cents=-106000,
            needs_review=True,
        ),
    ]
    summary = summarize(rows)
    assert summary.total_cents == 0
    target = tmp_path / "negative.xlsx"
    write_workbook(
        target,
        rows,
        summary,
        filters=ExportFilters(),
        generated_at="2026-10-28 12:00:00",
    )
    workbook = load_workbook(target)
    detail = workbook["发票明细"]
    assert float(detail["K2"].value) == 1060.0
    assert float(detail["K3"].value) == -1060.0

    categories = workbook["分类汇总"]
    assert categories["A2"].value == "差旅费"
    assert float(categories["E2"].value) == 1060.0
    assert categories["A3"].value == "其他费用"
    assert float(categories["E3"].value) == -1060.0
    # A zero net total cannot produce a meaningful share; show 0.0% instead
    # of dividing by zero.
    assert float(categories["F2"].value) == 0.0
    assert float(categories["F3"].value) == 0.0
    assert float(categories["E4"].value) == 0.0
    assert float(categories["F4"].value) == 0.0


def test_workbook_handles_rows_without_any_amounts(tmp_path: Path) -> None:
    rows = [
        make_row(
            id=1,
            source_path=str(tmp_path / "a.pdf"),
            invoice_type=InvoiceType.UNKNOWN.value,
            amount_net_cents=None,
            tax_amount_cents=None,
            amount_total_cents=None,
            confidence={"amount_total": "low"},
            needs_review=True,
        ),
        make_row(
            id=2,
            source_path=str(tmp_path / "b.pdf"),
            amount_net_cents=None,
            tax_amount_cents=None,
            amount_total_cents=None,
            needs_review=True,
        ),
    ]
    summary = summarize(rows)
    assert summary.missing_amount_count == 2
    assert summary.total_cents == 0

    target = tmp_path / "no-amounts.xlsx"
    write_workbook(
        target,
        rows,
        summary,
        filters=ExportFilters(),
        generated_at="2026-10-28 12:00:00",
    )
    workbook = load_workbook(target)
    detail = workbook["发票明细"]
    assert detail["D2"].value == "未识别"
    assert detail["I2"].value is None
    assert detail["J2"].value is None
    assert detail["K2"].value is None
    assert detail["K3"].value is None

    categories = workbook["分类汇总"]
    assert float(categories["E3"].value) == 0.0
    assert float(categories["F2"].value) == 0.0
    assert float(categories["F3"].value) == 0.0

    meta = {
        row[0].value: row[1].value
        for row in workbook["导出信息"].iter_rows(min_row=2)
        if row[0].value and row[1].value is not None
    }
    assert meta["金额缺失条数"] == 2
    assert float(meta["价税合计"]) == 0.0
